#!/usr/bin/env python3
"""Turn the source files into the openable ones.

    python3 tools/build_deliverables.py

    presentations/midterm-oct15.json  ->  presentations/ZenAI-Midterm-Oct15.pptx
    docs/build-plan.md               ->  docs/ZenAI-Build-Plan.docx
    docs/diagrams.md                 ->  docs/ZenAI-Diagrams-Source.docx

Edit the SOURCE (json / md), then re-run this. Never hand-edit the .pptx or
.docx -- they are build output and the next run overwrites them.

Needs: pip install python-pptx python-docx
"""
from __future__ import annotations

import json
import os
import re

from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ZenAI's palette -- same values as design/screens.html
PAPER = RGBColor(0xFB, 0xF9, 0xF6)
INK = RGBColor(0x1C, 0x26, 0x20)
MUTED = RGBColor(0x6B, 0x7A, 0x72)
GREEN = RGBColor(0x2F, 0x6F, 0x5E)
SOFT = RGBColor(0xE3, 0xEF, 0xE9)
LINE = RGBColor(0xE4, 0xE0, 0xD8)
SER = "Georgia"

TEAM = "Abraham Irabor  ·  Fnu Soh Tah Fon  ·  Joseph Williams  ·  Chinonso Egeolu"


def clean(t: str) -> str:
    """Strip markdown that has no meaning once rendered."""
    return re.sub(r"`([^`]+)`", r"\1", t or "")


def split_bold(text: str):
    """'a **b** c' -> [('a ',False),('b',True),(' c',False)]"""
    parts, buf, on = [], "", False
    i = 0
    while i < len(text):
        if text[i:i + 2] == "**":
            parts.append((buf, on)); buf = ""; on = not on; i += 2
        else:
            buf += text[i]; i += 1
    parts.append((buf, on))
    return [(c, b) for c, b in parts if c]


# --------------------------------------------------------------------- PPTX
# The deck's look comes from the ZenAI film itself: warm paper, deep green,
# heavy sans headlines, serif for anything ZenAI "says".
FILM_PAPER = RGBColor(0xF8, 0xF7, 0xF3)
CRISIS = RGBColor(0xA9, 0x3E, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD = "Helvetica Neue"
MONO = "Menlo"
SW, SH = 13.333, 7.5          # 16:9, inches

MEDIA_DIRS = ["", "presentations/media", "docs/diagrams"]
FILM = "presentations/media/zenai-film.mp4"


def _resolve(path: str):
    """Deck paths are workspace-relative; this repo keeps copies in
    presentations/media and docs/diagrams. Find whichever exists."""
    if not path:
        return None
    for d in MEDIA_DIRS:
        cand = os.path.join(d, path) if d == "" else os.path.join(d, os.path.basename(path))
        if os.path.exists(cand):
            return cand
    return None


def _runs(p, text, size, color, bold=False, font=HEAD, italic=False):
    """Inline markdown: **bold** and `code`."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text or ""):
        if not part:
            continue
        r = p.add_run()
        if part.startswith("**"):
            r.text, b, f = part[2:-2], True, font
        elif part.startswith("`"):
            r.text, b, f = part[1:-1], bold, MONO
        else:
            r.text, b, f = part, bold, font
        r.font.size = Pt(size if f != MONO else size * 0.9)
        r.font.color.rgb = color
        r.font.bold = b
        r.font.italic = italic
        r.font.name = f


def _box(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _text(slide, x, y, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
          font=HEAD, italic=False, anchor=MSO_ANCHOR.TOP, spacing=1.08):
    tf = _box(slide, x, y, w, h, anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    _runs(p, text, size, color, bold, font, italic)
    return tf


def _rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def _picture(slide, path, x, y, w, h, border=True):
    """Fit inside the box, keep aspect, centre."""
    from PIL import Image

    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    if border:
        _rect(slide, px - 0.04, py - 0.04, pw + 0.08, ph + 0.08, WHITE, LINE)
    slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))
    return px, py, pw, ph


def _title(slide, text, y=0.55):
    _text(slide, 0.75, y, 11.8, 0.9, text, 30, INK, bold=True)
    _rect(slide, 0.75, y + 0.88, 0.9, 0.05, GREEN)


def _footer(slide, n, total):
    _text(slide, 0.75, 7.02, 6, 0.3, "ZenAI  ·  COSC 490  ·  Group 3", 9.5, MUTED)
    _text(slide, SW - 2.75, 7.02, 2.0, 0.3, f"{n} / {total}", 9.5, MUTED, align=PP_ALIGN.RIGHT)


def _bullets(slide, x, y, w, h, items, size=18):
    tf = _box(slide, x, y, w, h)
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * 0.75)
        p.line_spacing = 1.1
        r = p.add_run()
        r.text = "●  "
        r.font.size = Pt(size * 0.55)
        r.font.color.rgb = GREEN
        r.font.name = HEAD
        _runs(p, b, size, INK)


def build_pptx(src="presentations/midterm-oct15.json",
               out="presentations/ZenAI-Midterm-Oct15.pptx"):
    deck = json.load(open(src))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]
    def has_content(sl):
        return bool(sl.get("title") or sl.get("body") or sl.get("image")
                    or any((b or "").strip() for b in sl.get("bullets", []))
                    or sl.get("objects"))

    # A slide with nothing on it (e.g. an accidental "add slide" in the editor)
    # is skipped rather than shipped blank. Original indices are kept so the
    # film slide is still found by position.
    slides = [(k, x) for k, x in enumerate(deck["slides"]) if has_content(x)]
    skipped = len(deck["slides"]) - len(slides)
    if skipped:
        print(f"   skipped {skipped} empty slide(s)")
    total = len(slides)
    film_done = False

    for n, (i, s) in enumerate(slides):
        lay = s.get("layout")
        sl = prs.slides.add_slide(blank)
        bg = GREEN if lay == "section" else FILM_PAPER
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = bg
        img = _resolve(s.get("image", ""))

        if lay == "title":
            # the film's logo lockup: green rounded mark + serif wordmark
            mark = sl.shapes.add_shape(5, Inches(4.05), Inches(2.72), Inches(1.15), Inches(1.15))
            mark.fill.solid(); mark.fill.fore_color.rgb = GREEN
            mark.line.fill.background(); mark.shadow.inherit = False
            mtf = mark.text_frame; mtf.vertical_anchor = MSO_ANCHOR.MIDDLE
            mp = mtf.paragraphs[0]; mp.alignment = PP_ALIGN.CENTER
            _runs(mp, "Z", 50, WHITE, font=SER)
            _text(sl, 5.45, 2.48, 5.5, 1.6, s["title"], 96, INK, font=SER,
                  anchor=MSO_ANCHOR.MIDDLE)
            _text(sl, 1.5, 4.35, 10.33, 0.7, s.get("subtitle", ""), 24, GREEN,
                  align=PP_ALIGN.CENTER)
            for o in s.get("objects", []):
                if o.get("kind") != "text":
                    continue
                col = GREEN if o.get("color") == "accent" else MUTED
                t = o["text"].upper() if o.get("uppercase") else o["text"]
                _text(sl, SW * o["x"] / 100, SH * o["y"] / 100, SW * o["w"] / 100,
                      SH * o["h"] / 100, t, o.get("size", 14), col,
                      bold=(o.get("weight", 400) >= 600), align=PP_ALIGN.CENTER)

        elif lay == "image" and i == 1 and os.path.exists(FILM) and not film_done:
            # THE FILM: a real embedded video, plays in PowerPoint and Keynote
            _title(sl, s.get("title", ""))
            vw = 8.8                      # leaves room for the caption above the footer
            vh = vw * 9 / 16
            vx, vy = (SW - vw) / 2, 1.7
            _rect(sl, vx - 0.05, vy - 0.05, vw + 0.1, vh + 0.1, WHITE, LINE)
            sl.shapes.add_movie(FILM, Inches(vx), Inches(vy), Inches(vw), Inches(vh),
                                poster_frame_image=img, mime_type="video/mp4")
            _text(sl, 0.75, vy + vh + 0.14, 11.8, 0.3,
                  "▶  Click to play  ·  62 seconds  ·  backup: zenai-screens.unv.run/film",
                  12.5, MUTED, align=PP_ALIGN.CENTER)
            film_done = True

        elif lay == "image":
            _title(sl, s.get("title", ""))
            if img:
                _picture(sl, img, 0.75, 1.72, 11.83, 4.75)
            if s.get("subtitle"):
                _text(sl, 0.75, 6.55, 11.83, 0.4, s["subtitle"], 12, MUTED,
                      align=PP_ALIGN.CENTER, italic=True)

        elif lay == "split":
            _title(sl, s.get("title", ""))
            _bullets(sl, 0.75, 1.85, 5.6, 4.9, s.get("bullets", []), size=18.5)
            if img:
                _picture(sl, img, 6.65, 1.8, 5.95, 4.9)

        elif lay == "bullets":
            _title(sl, s.get("title", ""))
            _bullets(sl, 0.75, 1.9, 11.8, 4.95, s.get("bullets", []), size=23)

        elif lay == "quote":
            _text(sl, 0.9, 1.35, 1.2, 1.2, "“", 120, GREEN, font=SER)
            _text(sl, 1.6, 2.15, 10.2, 2.6, s.get("body", ""), 42, INK, font=SER,
                  anchor=MSO_ANCHOR.MIDDLE, spacing=1.12)
            _rect(sl, 1.6, 5.0, 0.9, 0.05, GREEN)
            _text(sl, 1.6, 5.25, 10.2, 1.2, s.get("subtitle", ""), 15, MUTED)

        elif lay == "section":
            _text(sl, 1.0, 2.6, 11.3, 1.3, s.get("title", ""), 54, WHITE, bold=True)
            _text(sl, 1.0, 4.0, 11.3, 1.0, s.get("subtitle", ""), 20,
                  RGBColor(0xD6, 0xE8, 0xDF))

        elif lay == "stat":
            _text(sl, 0.75, 1.15, 11.83, 2.6, s.get("body", ""), 128, GREEN, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            _text(sl, 1.5, 3.95, 10.33, 0.8, s.get("title", ""), 30, INK, bold=True,
                  align=PP_ALIGN.CENTER)
            _text(sl, 2.0, 4.85, 9.33, 1.6, s.get("subtitle", ""), 17, MUTED,
                  align=PP_ALIGN.CENTER, spacing=1.2)

        if lay not in ("title", "section"):
            _footer(sl, n + 1, total)
        if s.get("note"):
            sl.notes_slide.notes_text_frame.text = s["note"]

    prs.save(out)
    return out


# --------------------------------------------------------------------- DOCX
DINK = DRGB(0x1C, 0x26, 0x20)
DGREEN = DRGB(0x2F, 0x6F, 0x5E)
DMUTED = DRGB(0x6B, 0x7A, 0x72)


def _dp(p, text, size=11, color=DINK, base_bold=False):
    for chunk, b in split_bold(clean(text)):
        r = p.add_run(chunk)
        r.font.size = DPt(size); r.font.color.rgb = color
        r.font.bold = b or base_bold; r.font.name = "Calibri"
    return p


def build_docx(md_path, out_path, title, subtitle):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = DInches(0.9)
        s.top_margin = s.bottom_margin = DInches(0.8)

    r = doc.add_paragraph().add_run(title)
    r.font.size = DPt(26); r.font.bold = True; r.font.color.rgb = DINK; r.font.name = "Calibri"
    r = doc.add_paragraph().add_run(subtitle)
    r.font.size = DPt(11); r.font.color.rgb = DGREEN; r.font.name = "Calibri"
    doc.add_paragraph()

    lines = open(md_path).read().split("\n")
    i = 0
    while i < len(lines):
        st = lines[i].strip()

        if st.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                i += 1
            if rows:
                cols = max(len(x) for x in rows)
                t = doc.add_table(rows=0, cols=cols); t.style = "Light Grid Accent 1"
                for ri, row in enumerate(rows):
                    cs = t.add_row().cells
                    for ci in range(cols):
                        _dp(cs[ci].paragraphs[0], row[ci] if ci < len(row) else "",
                            size=9.5, base_bold=(ri == 0))
                doc.add_paragraph()
            continue

        if st.startswith("```"):
            i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            r = doc.add_paragraph().add_run("\n".join(buf))
            r.font.name = "Menlo"; r.font.size = DPt(8.5); r.font.color.rgb = DMUTED
            continue

        if re.match(r"^#{1,4} ", st):
            lvl = len(st.split(" ")[0])
            p = doc.add_paragraph()
            p.paragraph_format.space_before = DPt(16 if lvl <= 2 else 10)
            p.paragraph_format.space_after = DPt(5)
            _dp(p, st[lvl + 1:], size={1: 20, 2: 16, 3: 13, 4: 11.5}[lvl],
                color=DINK if lvl <= 2 else DGREEN, base_bold=True)
        elif st.startswith("> "):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = DInches(0.35)
            _dp(p, st[2:], color=DGREEN)
        elif re.match(r"^[-*] ", st):
            _dp(doc.add_paragraph(style="List Bullet"), st[2:])
        elif re.match(r"^\d+\. ", st):
            _dp(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", st))
        elif st and not st.startswith("---"):
            p = doc.add_paragraph(); p.paragraph_format.space_after = DPt(7)
            _dp(p, st)
        i += 1

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    print("→", build_pptx())
    print("→", build_docx("docs/build-plan.md", "docs/ZenAI-Build-Plan.docx",
                          "ZenAI — Build Plan",
                          "COSC 490 · Project 5 · Group 3 · Sep 30 → Dec 3, 2026"))
    print("→", build_docx("docs/diagrams.md", "docs/ZenAI-Diagrams-Source.docx",
                          "ZenAI — Diagram source (Mermaid)",
                          "Paste into Lucidchart: Insert → Diagram as code → Mermaid"))
